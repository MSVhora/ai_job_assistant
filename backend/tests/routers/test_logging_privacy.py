import io
import json
import logging
import re

import pytest
from docx import Document
from fakes import VALID_PROFILE, FakeJobSource, fake_posting, install_acompletion, llm_response
from httpx import ASGITransport, AsyncClient

from app.adapters.job_sources import registry
from app.core.config import get_settings
from app.main import app

pytestmark = pytest.mark.usefixtures("clean_tables")

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
RESUME_MARKER = "ZEBRA-RESUME-MARKER-7421"
POSTING_MARKER = "NARWHAL-POSTING-MARKER-9913"
API_KEY = "AIzaSyTESTSENTINELKEY1234567890abcdefg"
KEY_SHAPED = re.compile(r"AIza[0-9A-Za-z_-]{20,}|Bearer\s+\S+|sk-[0-9A-Za-z]{20,}")


@pytest.fixture(autouse=True)
def gemini_key(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("GEMINI_API_KEY", API_KEY)
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


@pytest.fixture
async def client() -> AsyncClient:
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as async_client:
        yield async_client


def resume_docx() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph(f"Senior Data Analyst {RESUME_MARKER}")
    document.save(buffer)
    return buffer.getvalue()


async def test_resume_to_search_flow_logs_no_resume_text_prompts_or_keys(
    client: AsyncClient, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    caplog.set_level(logging.DEBUG)
    for name, existing in logging.root.manager.loggerDict.items():
        if name.startswith("app") and isinstance(existing, logging.Logger):
            existing.disabled = False  # alembic's fileConfig disables loggers created before it ran
    source = FakeJobSource("adzuna", postings=[fake_posting("p1", description=POSTING_MARKER)])
    monkeypatch.setattr(registry, "all_sources", lambda: (source,))
    prompts = install_acompletion(monkeypatch, lambda **kw: llm_response(json.dumps(VALID_PROFILE)))

    uploaded = (
        await client.post(
            "/api/resumes",
            files={"file": ("resume.docx", io.BytesIO(resume_docx()), DOCX_MIME)},
        )
    ).json()
    extracted = await client.post(f"/api/resumes/{uploaded['resume_id']}/extract")
    created = await client.post(
        "/api/profiles",
        json={
            "name": "Jane",
            "structured_profile": extracted.json()["draft_profile"],
            "source_resume_id": uploaded["resume_id"],
        },
    )
    search = await client.post(
        "/api/jobs/search",
        json={
            "query": "analyst",
            "profile_id": created.json()["profile_id"],
            "country": "de",
            "source": "adzuna",
        },
    )

    assert (extracted.status_code, created.status_code, search.status_code) == (200, 201, 202)
    assert prompts, "the flow must have exercised the LLM wrapper"
    logged = "\n".join(record.getMessage() for record in caplog.records)
    assert "task=" in logged
    assert logged
    assert RESUME_MARKER not in logged
    assert POSTING_MARKER not in logged
    assert API_KEY not in logged
    assert not KEY_SHAPED.search(logged)
    for call in prompts:
        for message in call["messages"]:
            assert message["content"][:80] not in logged


async def test_resume_document_flow_logs_no_profile_text_or_keys(
    client: AsyncClient, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    from app.adapters.evidence_sources.base import EvidenceSourceError
    from app.services import resume_documents

    class FailingGitHub:
        def is_configured(self) -> bool:
            return True

        async def identify(self) -> None:
            msg = f"boom {API_KEY}"
            raise EvidenceSourceError(msg)

    caplog.set_level(logging.DEBUG)
    for name, existing in logging.root.manager.loggerDict.items():
        if name.startswith("app") and isinstance(existing, logging.Logger):
            existing.disabled = False
    monkeypatch.setattr(resume_documents, "get_source", lambda _name: FailingGitHub())
    profile = {
        **VALID_PROFILE,
        "experience": [{"company": "Acme", "bullets": [f"Shipped {RESUME_MARKER}"]}],
    }
    created_profile = await client.post(
        "/api/profiles", json={"name": "Jane", "structured_profile": profile}
    )
    created = await client.post(
        "/api/resume-documents", json={"profile_id": created_profile.json()["profile_id"]}
    )
    document_id = created.json()["id"]
    conflicts = await client.get(f"/api/resume-documents/{document_id}/conflicts")
    exported = await client.get(f"/api/resume-documents/{document_id}/export")

    assert (created.status_code, conflicts.status_code, exported.status_code) == (201, 200, 200)
    assert conflicts.json()["github_checked"] is False
    assert RESUME_MARKER in exported.text
    logged = "\n".join(record.getMessage() for record in caplog.records)
    assert "resume.identity failed" in logged
    assert RESUME_MARKER not in logged
    assert API_KEY not in logged
    assert not KEY_SHAPED.search(logged)


async def test_resume_generation_flow_logs_no_jd_comment_prompt_or_evidence_text(
    client: AsyncClient, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    from fakes import FakeResumeLLM, seed_resume_world

    caplog.set_level(logging.DEBUG)
    for name, existing in logging.root.manager.loggerDict.items():
        if name.startswith("app") and isinstance(existing, logging.Logger):
            existing.disabled = False
    world = await seed_resume_world()
    llm = FakeResumeLLM(jd={"must_haves": [POSTING_MARKER], "keywords": ["Snowflake"]})
    prompts = install_acompletion(monkeypatch, llm)

    created = await client.post(
        "/api/resume-documents",
        json={
            "profile_id": str(world["profile"]),
            "job_description": f"We need {POSTING_MARKER} and Snowflake",
        },
    )
    document = created.json()
    entry = next(job for job in document["content"]["work"] if job["company"] == "Acme Corp")
    base = f"/api/resume-documents/{document['id']}"
    await client.post(
        f"{base}/comments",
        json={"target": {"section": "work", "block_id": entry["id"]}, "text": RESUME_MARKER},
    )
    applied = await client.post(f"{base}/apply-comments")
    edited = await client.patch(
        f"{base}/bullets/{entry['highlights'][0]['id']}", json={"text": f"Shipped {RESUME_MARKER}"}
    )

    assert (created.status_code, applied.status_code, edited.status_code) == (201, 200, 200)
    assert prompts, "the flow must have exercised the LLM wrapper"
    logged = "\n".join(record.getMessage() for record in caplog.records)
    assert "cost_usd=" in logged
    for secret in (RESUME_MARKER, POSTING_MARKER, "Faster nightly import", "Helm charts"):
        assert secret not in logged
    for call in prompts:
        for message in call["messages"]:
            assert message["content"][:80] not in logged
